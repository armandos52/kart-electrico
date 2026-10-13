from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT/'07_DESARROLLO/propuesta_2026-27/Proposta_Smart_Electric_Kart_V2_2026-27.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin = s.bottom_margin = Cm(1.8)
s.left_margin = s.right_margin = Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st = d.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_after = Pt(6)
d.styles['Normal'].font.size = Pt(10.5)
d.styles['Normal'].paragraph_format.line_spacing = 1.08
d.styles['Title'].font.size = Pt(23)
d.styles['Heading 1'].font.size = Pt(14)
d.styles['Heading 2'].font.size = Pt(11)
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
d.core_properties.title = 'Proposta de projeto Smart Electric Kart 2026 2027'
d.core_properties.author = 'Armando Aimacaña e Irvin Chasi'
s.header.paragraphs[0].text = 'SMART ELECTRIC KART  |  Proposta de projeto 2026/27'
s.header.paragraphs[0].runs[0].font.size = Pt(8)
f = s.footer.paragraphs[0]
f.text = 'V2.0 · 30 de setembro de 2026 · Proposta para discussão     |     '
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
f.runs[0].font.size = Pt(8)

def p(text, style=None): return d.add_paragraph(text, style)
def h(text): d.add_heading(text, 1)
def table(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = 'Table Grid'
    for c, text, w in zip(t.rows[0].cells, headers, widths):
        c.text=text; c.width=Cm(w)
        for r in c.paragraphs[0].runs: r.bold=True
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E9EEF2'); c._tc.get_or_add_tcPr().append(sh)
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=t.add_row().cells
        for c,text,w in zip(cells,row,widths): c.text=text; c.width=Cm(w)
    for row in t.rows:
        no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.space_before=Pt(3)
                for r in para.runs: r.font.size=Pt(9)
    spacer = d.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.4
    spacer.add_run().font.size = Pt(3)
    return t

p('Proposta de projeto Smart Electric Kart', 'Title')
p('Aperfeiçoamento do diferencial eletrónico e estimação da autonomia real', 'Subtitle')
p('Projeto em Engenharia Eletrotécnica e de Computadores · Ano letivo 2026/27')
p('Alunos: Armando Aimacaña e Irvin Chasi\nOrientadores indicados: Luís Neves e Paulo Ventura\nData: 30 de setembro de 2026 · Versão 2.0')
p('Propomos estudar a configuração recebida do kart, aperfeiçoar uma função do diferencial eletrónico e estimar a autonomia a partir de medições reproduzíveis. Até fevereiro de 2027, pretendemos apresentar uma implementação identificada, uma avaliação experimental inicial e uma estimativa de autonomia com condições e limitações explícitas. O âmbito e o calendário serão revistos com os orientadores.')
h('1 Enquadramento e problema')
p('O relatório de julho de 2026 descreve uma unidade frontal STM32, uma unidade traseira ESP32, comunicação CAN, BMS e dois motores BLDC traseiros com controladores independentes [1, p. PDF 5]. Esta arquitetura é a referência documental; a sua correspondência com o kart disponível ainda terá de ser verificada.')
p('O diferencial já possui modos DIF, BLQ e ESC. Os autores anteriores declaram ensaios de DIF e BLQ, mas não validaram experimentalmente ESC devido a uma falha na ligação GND do giroscópio GY-521 [1, p. PDF 97]. A contribuição proposta consiste em diagnosticar, aperfeiçoar e validar esta base. A função específica será selecionada após o diagnóstico, podendo incluir a validação e o aperfeiçoamento de ESC se os meios disponíveis o permitirem.')
p('O relatório também refere falhas do contactor, bloqueios do ESP32 e comportamento anómalo do SOC do BMS [1, pp. PDF 97–98]. A relação entre flutuações de alimentação e bloqueios é apresentada como hipótese. Estes problemas devem ser investigados antes de depender dos dados para avaliar o controlo ou a autonomia.')
h('2 Pergunta e objetivo geral')
p('Como aperfeiçoar e validar o diferencial eletrónico da configuração recebida e estimar a autonomia real do kart através de medições reproduzíveis?')
p('O objetivo geral é obter uma melhoria de controlo justificável por comparação experimental e uma estimativa de autonomia associada a um percurso e a condições de utilização definidos.')

d.add_page_break()
h('3 Objetivos específicos e resultados esperados')
table(['Objetivo','Resultado verificável'],[
('Identificar a configuração e os impedimentos iniciais','Inventário físico, versões de firmware, sinais disponíveis e registo de incidências.'),
('Definir a melhoria do diferencial','Função selecionada, justificação, indicadores e protocolo revistos com os orientadores.'),
('Garantir medições utilizáveis','Verificação do registo, calibrações e ensaio de referência com condições documentadas.'),
('Implementar e avaliar a melhoria','Código versionado e comparação entre referência e versão modificada.'),
('Estimar a autonomia até fevereiro de 2027','Dados de consumo e distância, método de cálculo e incerteza ou limitações identificadas.'),
('Melhorar a interface e consolidar resultados','Informação de utilização pertinente e validação final no segundo semestre.')],[6,11])
h('4 Metodologia e critérios de avaliação')
p('Começaremos pela leitura dirigida da documentação e pela inspeção do kart. Registaremos componentes, ligações, firmware, sinais e anomalias observadas. As correções necessárias à realização dos ensaios terão prioridade sobre a inclusão de novas funcionalidades. Novos sensores só serão propostos quando houver uma necessidade de medição demonstrada.')
p('Para o diferencial, estabeleceremos uma referência inicial e compararemos a versão modificada em condições equivalentes de percurso, carga e condução. Selecionaremos indicadores adequados à função escolhida, como erro de seguimento da velocidade das rodas, resposta em curva e erro de guinada quando aplicável. Os limiares de aceitação serão definidos antes dos ensaios comparativos, com base nas medições iniciais e na revisão com os orientadores.')
p('Para a autonomia, registaremos tempo, distância, tensão e corrente da bateria, estado inicial de carga e condições do percurso. Verificaremos unidades, sincronização e qualidade dos dados. Quando as medições forem adequadas, calcularemos energia em Wh e consumo em Wh/km. A autonomia estimada em km será obtida pela relação entre energia utilizável medida em Wh e consumo em Wh/km, para as condições ensaiadas. Não assumiremos que o SOC indicado pelo BMS, isoladamente, determina a energia utilizável.')
p('Documentaremos a precisão conhecida dos instrumentos, a dispersão entre repetições e os dados em falta. Distinguiremos a autonomia estimada da autonomia validada num percurso. Se não for possível medir a energia utilizável até fevereiro, apresentaremos o consumo medido e uma estimativa condicionada, identificando a evidência ainda necessária.')
h('5 Âmbito e dependências')
p('O núcleo do trabalho inclui diagnóstico, qualidade de dados, melhoria selecionada do diferencial e avaliação de consumo e autonomia. A interface será aperfeiçoada em março e abril, de acordo com os resultados. Alterações extensas ao sistema de potência ou novas funções de controlo exigirão revisão do âmbito, dos recursos e do calendário.')

d.add_page_break()
h('6 Calendarização proposta')
p('Planeamento de outubro de 2026 a junho de 2027. As datas de avaliação académica permanecem por confirmar. Os ensaios decorrem ao longo do desenvolvimento; a documentação acompanha todas as fases.')
headers=['Atividade','Out','Nov','Dez','Jan','Fev','Mar','Abr','Mai','Jun']
activities=[
('Diagnóstico e estudo', [1,0,0,0,0,0,0,0,0]),
('Âmbito e protocolo', [1,1,0,0,0,0,0,0,0]),
('Correções e calibração', [0,1,1,0,0,0,0,0,0]),
('Ensaios de referência', [0,1,1,0,0,0,0,0,0]),
('Melhoria do controlo', [0,0,1,1,1,0,0,0,0]),
('Consumo e autonomia', [0,0,0,1,1,0,0,0,0]),
('Interface de utilização', [0,0,0,0,0,1,1,0,0]),
('Validação final', [0,0,0,0,0,0,0,1,1]),
('Documentação contínua', [1]*9),
('Revisão e correções', [0,0,0,0,1,0,1,1,1])]
table(headers, [(name,*['X' if x else '' for x in flags]) for name,flags in activities],[6.2]+[1.2]*9)
h('7 Marcos e entregáveis')
table(['Prazo proposto','Entregável e condição de passagem'],[
('Fim de outubro','M1 — Configuração de partida documentada e diagnóstico dos impedimentos aos ensaios.'),
('Fim de novembro','M2 — Função de controlo, indicadores, recursos e protocolo definidos; responsabilidades acordadas.'),
('Fim de dezembro','M3 — Medições de referência utilizáveis e versão inicial da melhoria, se o diagnóstico o permitir.'),
('Fim de fevereiro','M4 — Relatório intercalar: implementação, comparação inicial e estimativa de autonomia com condições, incerteza e limites.'),
('Fim de abril','M5 — Interface aperfeiçoada e verificação das funções afetadas pelas alterações.'),
('Fim de junho','M6 — Validação final, relatório, código e dados organizados para entrega; ajustar à data oficial.')],[3.4,13.6])
p('O marco de fevereiro depende da disponibilidade do kart e da resolução dos impedimentos de funcionamento e medição. Se estes persistirem, o relatório intercalar deverá apresentar o diagnóstico sustentado, os resultados obtidos e o plano de conclusão revisto.')

d.add_page_break()
h('8 Organização e recursos')
p('Armando Aimacaña e Irvin Chasi desenvolverão o trabalho com acompanhamento dos orientadores indicados. A distribuição individual das tarefas será acordada no início e registada no diário do projeto. Propomos uma revisão no encerramento de cada marco, a agendar com os orientadores.')
p('Recursos a confirmar: acesso ao kart e à configuração de programação correspondente, instrumentos de medição elétrica, aquisição ou registo sincronizado de dados e local autorizado para ensaios. A escolha de novos componentes dependerá do diagnóstico, da compatibilidade e dos recursos disponíveis.')
h('9 Riscos e resposta prevista')
table(['Risco ou dependência','Resposta proposta'],[
('Falhas de alimentação ou controlo','Diagnosticar antes dos ensaios integrados e conservar evidência das correções.'),
('Sensores ou registos inconsistentes','Verificar firmware, unidades, calibração e sincronização; repetir medições inválidas.'),
('SOC ou energia utilizável sem validação','Não usar o SOC como única evidência; declarar limites da estimativa de autonomia.'),
('Indisponibilidade do veículo ou atraso de componentes','Priorizar o núcleo do projeto e rever tarefas opcionais e marcos.'),
('Melhoria sem benefício demonstrado','Registar o resultado comparativo e rever a estratégia, sem declarar sucesso sem evidência.')],[6,11])
h('10 Documentação e entrega')
p('Manteremos um diário desde outubro, com decisões, configuração utilizada, versões de software, calibrações e condições de cada ensaio. Cada conjunto de dados será ligado ao firmware e ao protocolo que o produziu. Os resultados serão apresentados com unidades, condições e referências às fontes; os originais históricos serão preservados.')
p('A entrega deverá reunir relatório, configuração e código identificados, dados de referência e comparação, método de estimação de autonomia e lista de limitações ou trabalho futuro. A proposta será revista após o diagnóstico inicial para fixar a função do diferencial, os critérios quantitativos e a distribuição de responsabilidades.')
h('11 Referências e base documental')
p('[1] Relatório de projeto Kart Cross, Bernardo 2223403 e Ricardo 2223406, julho de 2026. Catálogo DOC-0042. Referências utilizadas: páginas PDF 5 e 96–98 (conclusões, páginas impressas 82–84). Descreve a configuração de 2025/26; montagem atual por confirmar.')
p('[2] Cronograma Projeto Smart Electric Kart V1.0, datado de 30 de setembro de 2026. Catálogo DOC-0043. Base desta proposta: tabela de tarefas e objetivos para fevereiro de 2027.')
p('Esta versão é uma proposta para discussão. A seleção da função de controlo, os recursos, os limiares de aceitação e as datas académicas aguardam definição com os orientadores.')
d.save(OUT)
print(OUT)
